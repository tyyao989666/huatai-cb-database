"""Read cached values from the weekly workbook; never modifies the source Excel."""
from pathlib import Path
import argparse
import csv
import json
import re
import zipfile
from datetime import datetime
import openpyxl
from stock_names import (
    CNINFO_NAMES_URL, clean_stock_name, fetch_cninfo_names,
    normalize_stock_code, read_name_cache, resolve_stock_names,
)

BASE = Path(__file__).parent

def write_csv(name, fields, rows):
    with (BASE / 'data' / name).open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)

def update(source, reference=None, offline_stock_names=False):
    w = openpyxl.load_workbook(source, read_only=True, data_only=True)
    name_cache_path = BASE / 'data/stock_names.json'
    old_names = read_name_cache(name_cache_path)
    with (BASE / 'data/bonds.csv').open(encoding='utf-8-sig') as f:
        for row in csv.DictReader(f):
            code = normalize_stock_code(row.get('stock_code'))
            name = clean_stock_name(row.get('stock_name'))
            if code and name:
                old_names[code] = name
    fields = 'code name industry stock_code price stock_price change stock_change remaining turnover convert_price parity premium ytm floor floor_premium rating balance issue_size change_5d stock_change_5d list_date'.split()
    bonds = []
    for row in w['转债个券数据'].iter_rows(min_row=20, max_col=22, values_only=True):
        if not isinstance(row[0], str) or not re.fullmatch(r'\d{6}\.(SH|SZ)', row[0]):
            continue
        item = dict(zip(fields, row))
        item['stock_name'] = old_names.get(item['stock_code'], '')
        bonds.append(item)
    all_names, fetched_names, name_warning = resolve_stock_names(
        bonds, old_names, fetcher=None if offline_stock_names else fetch_cninfo_names,
    )
    market_fields = 'date implied_vol implied_vol_smooth implied_vol_percentile price_median turnover_total parity_median premium_median floor_premium_median floor_median remaining_median ytm_median'.split()
    market = []
    for row in w['转债市场数据'].iter_rows(min_row=11, max_col=12, values_only=True):
        if isinstance(row[0], datetime):
            item = dict(zip(market_fields, row))
            item['date'] = row[0].date().isoformat()
            market.append(item)
    supply = []
    for row in w['转债供给'].iter_rows(min_row=13, max_col=3, values_only=True):
        if isinstance(row[0], datetime):
            supply.append(dict(date=row[0].date().isoformat(), balance=row[1], count=row[2]))
    quantiles = []
    for row in w['转债估值'].iter_rows(min_row=12, max_row=16, min_col=2, max_col=7, values_only=True):
        quantiles.append(dict(zip(['window','implied_vol','parity','price','ytm','premium'], row)))
    listings = []
    for row in w['附1 _转债上市数据'].iter_rows(min_row=35, max_col=11, values_only=True):
        if isinstance(row[1], str) and re.fullmatch(r'\d{6}\.(SH|SZ)', row[1]):
            listings.append(dict(zip(['date','code','name','size','first_price','first_parity','first_premium','rating','first_turnover','industry','source_change'], row)))
    pending = []
    for row in w['附1 _转债上市数据'].iter_rows(min_row=11, max_row=30, max_col=6, values_only=True):
        if isinstance(row[1], str) and re.fullmatch(r'\d{6}\.(SH|SZ)', row[1]):
            pending.append(dict(zip(['date','code','name','size','industry','rating'], row)))
    assert bonds and market and supply
    assert len({b['code'] for b in bonds}) == len(bonds), 'Duplicate bond codes'
    if fetched_names:
        cache = json.loads(name_cache_path.read_text(encoding='utf-8')) if name_cache_path.exists() else {}
        cache['names'] = dict(sorted(all_names.items()))
        cache.setdefault('sources', []).append({
            'retrieved_at': datetime.now().astimezone().isoformat(timespec='seconds'),
            'url': CNINFO_NAMES_URL,
            'codes': sorted(fetched_names),
        })
        name_cache_path.write_text(json.dumps(cache, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    for name, cols, rows in [('bonds.csv',fields+['stock_name'],bonds),('market_history.csv',market_fields,sorted(market,key=lambda x:x['date'])),('supply_history.csv',['date','balance','count'],sorted(supply,key=lambda x:x['date']))]:
        write_csv(name, cols, rows)
    date = max(r['date'] for r in market)
    meta = dict(data_date=date, source=Path(source).name, bond_count=len(bonds), stock_names_missing=sum(not b['stock_name'] for b in bonds), quantiles=quantiles, listings=listings, pending=pending, reference='htscficc-analysis：仅借鉴视觉布局，未混入宏观指标数据')
    meta['stock_name_missing_codes'] = sorted({b['stock_code'] for b in bonds if not b['stock_name']})
    meta['stock_name_lookup_warning'] = name_warning
    (BASE / 'data/metadata.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    if reference:
        with zipfile.ZipFile(reference) as z:
            logo = z.read('static/images/huatai-logo.png')
        (BASE / 'assets').mkdir(exist_ok=True)
        (BASE / 'assets/huatai-logo.png').write_bytes(logo)
    w.close()
    print(json.dumps({k:meta[k] for k in ['data_date','bond_count','stock_names_missing']},ensure_ascii=False))
    if name_warning or meta['stock_name_missing_codes']:
        print(name_warning or '以下正股名称待补充：', meta['stock_name_missing_codes'])

if __name__ == '__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('source')
    parser.add_argument('--reference')
    parser.add_argument('--offline-stock-names', action='store_true', help='仅使用已核实的本地正股名称，不访问网络')
    args=parser.parse_args()
    update(args.source,args.reference,args.offline_stock_names)
